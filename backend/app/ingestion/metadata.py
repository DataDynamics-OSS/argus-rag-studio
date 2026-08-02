# SPDX-License-Identifier: Apache-2.0
"""문서 메타데이터 추출 — 본문(텍스트)과 별개로 문서 속성(제목/작성자/일시 등)을 뽑는다.

파서(``parsers.py``)가 본문을 추출한다면, 여기서는 파일 포맷별 *메타데이터*(문서 속성)를
추출해 ``rag_documents.metadata_json`` 에 저장한다. 검색 필터·출처 표시·정렬에 활용한다.

best-effort 원칙: 어떤 경우에도 예외를 던지지 않는다(메타데이터 부재가 인제스천을 막지 않음).
지원: hwp(OLE) / hwpx(OWPML) / pdf / docx / xlsx / pptx / doc·ppt·xls(OLE). 그 외는 빈 dict.

정규화 키(있을 때만): ``format`` ``title`` ``author`` ``last_saved_by`` ``created``
``modified`` ``app`` ``pages`` ``sections`` (+ 포맷별 extras).
"""

import logging
import struct
import zipfile
from xml.etree import ElementTree as ET

logger = logging.getLogger(__name__)


def _ext(filename: str) -> str:
    name = (filename or "").lower()
    dot = name.rfind(".")
    return name[dot:] if dot >= 0 else ""


# 파일 내부 producer 코드 → 사람이 읽는 앱 이름(매핑표).
# HWP/HWPX 는 실제 앱이 "아래아 한글" 이지만 파일에는 내부 코드로 기록된다:
#   - HWP(OLE): FileHeader 의 바이너리 포맷 버전(예: 5.0.5.0)
#   - HWPX(OWPML): version.xml 의 targetApplication = "WORDPROCESSOR"
# 둘 다 아래아 한글 네이티브 포맷이므로 앱 이름을 "아래아 한글" 로 정규화한다.
HANCOM_APP = "아래아 한글"
_APP_NAMES = {"WORDPROCESSOR": HANCOM_APP}

# ── HWP v5 (OLE / MS-OLEPS 속성셋) ───────────────────────────────────────────
# HwpSummaryInformation PropertyID(PIDSI) → 정규화 키
_PIDSI = {2: "title", 3: "subject", 4: "author", 5: "keywords",
          6: "comments", 8: "last_saved_by", 12: "created", 13: "modified", 18: "app"}


def _filetime_to_iso(raw8: bytes) -> str:
    import datetime

    v = struct.unpack("<Q", raw8)[0]
    if not v:
        return ""
    return (datetime.datetime(1601, 1, 1) + datetime.timedelta(microseconds=v // 10)).strftime(
        "%Y-%m-%d %H:%M:%S"
    )


def _parse_ole_propset(raw: bytes) -> dict:
    """MS-OLEPS 속성셋(첫 섹션)을 {pid: value} 로 파싱(VT_LPWSTR/LPSTR/FILETIME/I4)."""
    off0 = struct.unpack("<I", raw[44:48])[0]  # FMTID0(16) 뒤의 offset
    ps = raw[off0:]
    num = struct.unpack("<I", ps[4:8])[0]
    props: dict = {}
    for i in range(num):
        pid, poff = struct.unpack("<II", ps[8 + i * 8 : 16 + i * 8])
        vt = struct.unpack("<H", ps[poff : poff + 2])[0]
        body = ps[poff + 4 :]
        if vt == 0x1F:  # VT_LPWSTR (UTF-16)
            ln = struct.unpack("<I", body[:4])[0]
            props[pid] = body[4 : 4 + ln * 2].decode("utf-16le", "replace").rstrip("\x00")
        elif vt == 0x1E:  # VT_LPSTR
            ln = struct.unpack("<I", body[:4])[0]
            props[pid] = body[4 : 4 + ln].decode("cp949", "replace").rstrip("\x00")
        elif vt == 0x40:  # VT_FILETIME
            props[pid] = _filetime_to_iso(body[:8])
        elif vt == 0x03:  # VT_I4
            props[pid] = struct.unpack("<i", body[:4])[0]
    return props


def _meta_hwp(data: bytes) -> dict:
    import olefile

    out = {"format": "hwp"}
    ole = olefile.OleFileIO(__import__("io").BytesIO(data))
    try:
        fh = ole.openstream("FileHeader").read()
        v = struct.unpack("<I", fh[32:36])[0]  # 바이너리 포맷 버전
        out["app"] = f"{HANCOM_APP} {(v>>24)&0xff}.{(v>>16)&0xff}.{(v>>8)&0xff}.{v&0xff}"
        flags = struct.unpack("<I", fh[36:40])[0]
        fl = [n for b, n in [(1, "compressed"), (2, "password"), (4, "distribution"),
                             (8, "script"), (16, "drm")] if flags & b]
        if fl:
            out["flags"] = fl
        out["sections"] = sum(1 for s in ole.listdir() if s and s[0] == "BodyText")
        if ole.exists("\x05HwpSummaryInformation"):
            props = _parse_ole_propset(ole.openstream("\x05HwpSummaryInformation").read())
            for pid, key in _PIDSI.items():
                val = props.get(pid)
                if val not in (None, "", 0):
                    out[key] = str(val)[:300]
    finally:
        ole.close()
    return out


# ── HWPX (OWPML / ZIP) ───────────────────────────────────────────────────────
def _meta_hwpx(data: bytes) -> dict:
    out = {"format": "hwpx"}
    with zipfile.ZipFile(__import__("io").BytesIO(data)) as z:
        names = z.namelist()
        try:
            a = ET.fromstring(z.read("version.xml")).attrib
            # targetApplication 은 "WORDPROCESSOR" 같은 내부 코드 → 매핑표로 앱 이름 정규화(버전 유지).
            app_code = a.get("tagetApplication") or a.get("targetApplication") or ""
            app_name = _APP_NAMES.get(app_code, app_code) or HANCOM_APP
            ver = ".".join(filter(None, [a.get("major"), a.get("minor"), a.get("micro"), a.get("buildNumber")]))
            out["app"] = f"{app_name} {ver}".strip()
        except Exception:
            pass
        hpf = next((n for n in names if n.endswith("content.hpf")), None)
        if hpf:
            # content.hpf 의 OPF 메타데이터: <opf:title>, <opf:meta name=".."> 값은 엘리먼트 텍스트
            key_map = {"creator": "author", "lastsaveby": "last_saved_by",
                       "CreatedDate": "created", "ModifiedDate": "modified",
                       "subject": "subject", "keyword": "keywords", "description": "comments"}
            try:
                for el in ET.fromstring(z.read(hpf)).iter():
                    tag = el.tag.split("}")[-1]
                    txt = (el.text or "").strip()
                    if not txt:
                        continue
                    if tag == "title":
                        out["title"] = txt[:300]
                    elif tag == "language":
                        out["language"] = txt[:20]
                    elif tag == "meta":
                        nm = el.attrib.get("name")
                        if nm in key_map:
                            out[key_map[nm]] = txt[:300]
            except Exception:
                pass
        out["sections"] = sum(1 for n in names if n.startswith("Contents/section"))
    return out


# ── PDF / DOCX / XLSX (보너스 — 라이브러리 내장 속성) ─────────────────────────
def _meta_pdf(data: bytes) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(__import__("io").BytesIO(data))
    out = {"format": "pdf", "pages": len(reader.pages)}
    info = reader.metadata or {}
    for src, key in [("/Title", "title"), ("/Author", "author"), ("/Subject", "subject"),
                     ("/Creator", "app"), ("/CreationDate", "created"), ("/ModDate", "modified")]:
        val = info.get(src)
        if val:
            out[key] = str(val)[:300]
    return out


def _meta_docx(data: bytes) -> dict:
    import docx

    cp = docx.Document(__import__("io").BytesIO(data)).core_properties
    out = {"format": "docx"}
    for attr, key in [("title", "title"), ("author", "author"), ("subject", "subject"),
                      ("created", "created"), ("modified", "modified"), ("last_modified_by", "last_saved_by")]:
        val = getattr(cp, attr, None)
        if val:
            out[key] = str(val)[:300]
    return out


def _meta_xlsx(data: bytes) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(__import__("io").BytesIO(data), read_only=True)
    p = wb.properties
    out = {"format": "xlsx", "sheets": len(wb.sheetnames)}
    for attr, key in [("title", "title"), ("creator", "author"), ("subject", "subject"),
                      ("created", "created"), ("modified", "modified"), ("lastModifiedBy", "last_saved_by")]:
        val = getattr(p, attr, None)
        if val:
            out[key] = str(val)[:300]
    wb.close()
    return out


def _meta_ole_office(data: bytes, fmt: str) -> dict:
    """구형 오피스(.doc/.ppt/.xls, OLE2)의 ``\x05SummaryInformation`` 속성셋 추출.

    HWP 와 동일한 MS-OLEPS PIDSI 포맷이라 ``_parse_ole_propset``/``_PIDSI`` 를 그대로 재사용한다.
    """
    import olefile

    out = {"format": fmt}
    ole = olefile.OleFileIO(__import__("io").BytesIO(data))
    try:
        if ole.exists("\x05SummaryInformation"):
            props = _parse_ole_propset(ole.openstream("\x05SummaryInformation").read())
            for pid, key in _PIDSI.items():
                val = props.get(pid)
                if val not in (None, "", 0):
                    out[key] = str(val)[:300]
    finally:
        ole.close()
    return out


def _meta_pptx(data: bytes) -> dict:
    """PPTX(OOXML/ZIP) — docProps/core.xml(Dublin Core) + app.xml(슬라이드 수·앱) 직접 파싱.

    docx/xlsx 와 동일한 OOXML 구조라 라이브러리 없이 표준 메타를 뽑는다.
    """
    out = {"format": "pptx"}
    core_map = {
        "title": "title", "creator": "author", "subject": "subject",
        "description": "comments", "keywords": "keywords",
        "lastModifiedBy": "last_saved_by", "created": "created", "modified": "modified",
    }
    with zipfile.ZipFile(__import__("io").BytesIO(data)) as z:
        names = set(z.namelist())
        if "docProps/core.xml" in names:
            try:
                for el in ET.fromstring(z.read("docProps/core.xml")):
                    tag = el.tag.split("}")[-1]
                    txt = (el.text or "").strip()
                    if txt and tag in core_map:
                        out[core_map[tag]] = txt[:300]
            except Exception:
                pass
        if "docProps/app.xml" in names:
            try:
                for el in ET.fromstring(z.read("docProps/app.xml")):
                    tag = el.tag.split("}")[-1]
                    txt = (el.text or "").strip()
                    if not txt:
                        continue
                    if tag == "Slides":
                        try:
                            out["slides"] = int(txt)
                        except ValueError:
                            pass
                    elif tag == "Application":
                        out["app"] = txt[:300]
            except Exception:
                pass
    return out


_EXTRACTORS = {
    ".hwp": _meta_hwp,
    ".hwpx": _meta_hwpx,
    ".pdf": _meta_pdf,
    ".docx": _meta_docx,
    ".xlsx": _meta_xlsx,
    ".pptx": _meta_pptx,
    ".doc": lambda d: _meta_ole_office(d, "doc"),
    ".ppt": lambda d: _meta_ole_office(d, "ppt"),
    ".xls": lambda d: _meta_ole_office(d, "xls"),
}


def extract_metadata(filename: str, data: bytes) -> dict:
    """파일에서 문서 메타데이터를 추출한다(best-effort, 절대 예외 없음).

    지원 포맷이 아니거나 추출에 실패하면 빈 dict 를 반환한다. 빈 문자열 값은 제거한다.
    """
    fn = _EXTRACTORS.get(_ext(filename))
    if fn is None:
        return {}
    try:
        meta = fn(data)
    except Exception as e:  # 메타데이터 실패는 인제스천을 막지 않는다
        logger.warning("메타데이터 추출 실패(file=%s): %s", filename, e)
        return {}
    return {k: v for k, v in meta.items() if v not in (None, "", [])}
